"""Export DTECN-LLM final metrics to a 27-case Excel table and CSV.

Use recovery_v1/icl_final_test.json for recovered runs. Input files are read-only.
"""
import argparse
import csv
import json
import math
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side


METRICS = (
    "deadline_violation_rate",
    "max_fuzzy_lateness",
    "mean_fuzzy_lateness",
    "fuzzy_energy_score",
)
SCENARIOS = ("SS", "SM", "SL", "MS", "MM", "ML", "LS", "LM", "LL")
CASE_ORDER = tuple(f"{scenario}_{ddl}" for scenario in SCENARIOS for ddl in ("T", "M", "L"))


def collect_results(paths):
    results = {}
    for path in paths:
        data = json.loads(Path(path).read_text(encoding="utf-8"))
        if data.get("split") != "final_test":
            raise ValueError(f"Not a final-test result: {path}")
        teacher = data["identity"]["case_identity"]["teacher_identity"]
        if teacher.get("method") != "dtecn_llm":
            raise ValueError(f"Not a DTECN-LLM result: {path}")
        protocol = teacher["protocol"]
        ddl = protocol["ddl_setting"][0].upper()
        groups = data["scenario_results"]
        if not groups:
            raise ValueError(f"Empty scenario results: {path}")
        for scenario, group in groups.items():
            case = f"{scenario}_{ddl}"
            if case not in CASE_ORDER:
                raise ValueError(f"Unknown scenario/deadline: {case}")
            seeds = [int(record["seed"]) for record in group["seed_records"]]
            if len(seeds) != 30 or set(seeds) != set(range(201, 231)):
                raise ValueError(f"Expected exactly seeds 201-230: {case}")
            if case in results:
                raise ValueError(f"Duplicate case; select one final result: {case}")
            values = {}
            for metric in METRICS:
                value = float(group["aggregate"][metric])
                if not math.isfinite(value):
                    raise ValueError(f"Non-finite metric: {case}/{metric}")
                values[metric] = value
            results[case] = values
    return results


def export_table(paths, output):
    # Validate all inputs before writing either output.
    results = collect_results(paths)
    output = Path(output)
    if output.suffix.lower() != ".xlsx":
        raise ValueError("Output must have an .xlsx suffix")
    if output.resolve() in {Path(p).resolve() for p in paths}:
        raise ValueError("Output must not overwrite an input")
    wb = Workbook()
    ws = wb.active
    ws.title = "DTECN-LLM"
    ws.merge_cells("A1:A2")
    ws["A1"] = "场景"
    for col, metric in enumerate(METRICS, 2):
        ws.cell(1, col, metric)
        ws.cell(2, col, "DTECN-LLM")
    rows = []
    for row_index, case in enumerate(CASE_ORDER, 3):
        values = results.get(case, {})
        ws.cell(row_index, 1, case)
        for col, metric in enumerate(METRICS, 2):
            cell = ws.cell(row_index, col, values.get(metric))
            cell.number_format = "0.000"
        rows.append({"scenario": case, **values})
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows():
        for cell in row:
            cell.font = Font(name="Times New Roman", size=11)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border
    ws.column_dimensions["A"].width = 12
    for col in "BCDE":
        ws.column_dimensions[col].width = 27
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 22
    ws.freeze_panes = "B3"
    ws.sheet_view.showGridLines = False
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    csv_path = output.with_suffix(".csv")
    with csv_path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=["scenario", *METRICS])
        writer.writeheader()
        writer.writerows(rows)
    print(f"Excel: {output.resolve()}")
    print(f"CSV: {csv_path.resolve()}")
    missing = [case for case in CASE_ORDER if case not in results]
    if missing:
        print("Missing cases left blank: " + ", ".join(missing))


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, nargs="+", required=True,
                        help="Final icl_final_test.json files; select only one run per case")
    parser.add_argument("--output", type=Path, default=Path("DTECN_LLM.xlsx"))
    args = parser.parse_args(argv)
    export_table(args.input, args.output)


if __name__ == "__main__":
    main()
