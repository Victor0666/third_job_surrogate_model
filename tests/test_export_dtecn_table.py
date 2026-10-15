import csv
import json

from openpyxl import load_workbook
import pytest

from tools.export_dtecn_table import METRICS, collect_results, export_table


def result_file(tmp_path):
    group = dict(seed_records=[{"seed": s} for s in range(201, 231)],
                 aggregate=dict(zip(METRICS, (0.0123456, 200.1, 4.2, 333444.56789))))
    data = dict(split="final_test", identity={"case_identity": {"teacher_identity": {
        "method": "dtecn_llm", "protocol": {"ddl_setting": "tight"}}}},
        scenario_results={"SS": group})
    path = tmp_path / "icl_final_test.json"
    path.write_text(json.dumps(data), encoding="utf-8")
    return path, data


def test_layout_missing_cells_and_csv_precision(tmp_path):
    path, data = result_file(tmp_path)
    original = path.read_bytes()
    output = tmp_path / "DTECN_LLM.xlsx"
    export_table([path], output)
    wb = load_workbook(output)
    ws = wb.active
    assert ws.max_row == 29 and ws.max_column == 5
    assert ws["B1"].value == METRICS[0]
    assert ws["B2"].value == "DTECN-LLM"
    assert ws["A3"].value == "SS_T" and ws["A29"].value == "LL_L"
    assert ws["B3"].value == 0.0123456
    assert ws["B4"].value is None
    assert ws["B3"].number_format == "0.000"
    assert str(next(iter(ws.merged_cells.ranges))) == "A1:A2"
    assert ws.freeze_panes == "B3"
    wb.close()
    with output.with_suffix(".csv").open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    assert len(rows) == 27
    assert rows[0][METRICS[-1]] == "333444.56789"
    assert rows[1][METRICS[0]] == ""
    assert path.read_bytes() == original


@pytest.mark.parametrize("invalid", ["duplicate_seed", "missing_seed", "validation", "nan", "method"])
def test_invalid_results_rejected_before_output(tmp_path, invalid):
    path, data = result_file(tmp_path)
    group = data["scenario_results"]["SS"]
    if invalid == "duplicate_seed":
        group["seed_records"][-1]["seed"] = 201
    elif invalid == "missing_seed":
        group["seed_records"].pop()
    elif invalid == "validation":
        data["split"] = "validation"
    elif invalid == "method":
        data["identity"]["case_identity"]["teacher_identity"]["method"] = "other"
    else:
        group["aggregate"][METRICS[0]] = float("nan")
    path.write_text(json.dumps(data))
    output = tmp_path / "bad.xlsx"
    with pytest.raises(ValueError):
        export_table([path], output)
    assert not output.exists()


def test_duplicate_runs_are_not_silently_overwritten(tmp_path):
    path, _ = result_file(tmp_path)
    with pytest.raises(ValueError, match="Duplicate"):
        collect_results([path, path])
